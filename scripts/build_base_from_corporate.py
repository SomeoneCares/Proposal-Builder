#!/usr/bin/env python3
"""Build templates/vertowave_base.docx from Verto Wave's Word template.

Every proposal starts from that base, so it carries the corporate cover page,
styles, theme, running header and footer. Re-run this when the corporate
template is reissued:

    python scripts/build_base_from_corporate.py "Verto Wave Engagement Technical Proposal Template_V2.00.docx"

What it does to the supplied file:

1. Keeps the cover page and drops every placeholder block after it, including
   the table of contents — the builder writes its own.
2. Removes the reviewer comments that ship inside the template, so they cannot
   reach a customer.
3. Adds an empty first-page header. The Header style carries a bottom border,
   and without this the cover page shows that rule on its own.
4. Lifts the logo anchored in the running header clear of that same rule.
"""

from __future__ import annotations

import argparse
import os
import re
import shutil
import sys
import zipfile
from pathlib import Path

from docx import Document

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
COMMENT_PARTS = ("comments.xml", "commentsExtended.xml", "commentsIds.xml", "commentsExtensible.xml", "people.xml")
# The logo is anchored at the paragraph top; this lifts it above the Header style's rule.
LOGO_LIFT_EMU = -228600  # 0.25 inch
FIRST_PAGE_HEADER = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:p><w:pPr><w:pStyle w:val="Header"/><w:pBdr><w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/></w:pBdr></w:pPr></w:p>
</w:hdr>"""


def trim_to_cover(src: Path, dest: Path) -> int:
    doc = Document(str(src))
    body = doc.element.body
    children = list(body)
    cut = None
    for i, child in enumerate(children):
        if child.tag != W + "p":
            continue
        style = child.find(f"{W}pPr/{W}pStyle")
        text = "".join(t.text or "" for t in child.iter(W + "t"))
        if (style is not None and style.get(W + "val") in ("TOC", "TOCHeading")) \
                or text.strip().startswith("Table of Contents"):
            cut = i
            break
    if not cut:
        raise SystemExit("ERROR: could not find where the cover ends (no table-of-contents heading)")
    for child in children[cut:]:
        if child.tag != W + "sectPr":
            body.remove(child)
    doc.save(str(dest))
    return len(children) - cut - 1


def rewrite(src: Path, dest: Path) -> None:
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        names = set(zin.namelist())
        first_header = "word/headerFirst.xml"
        rel_id = "rIdFirstHeader"
        for item in zin.infolist():
            name = item.filename
            if any(name == "word/" + part for part in COMMENT_PARTS):
                continue
            data = zin.read(name)
            if name == "word/document.xml":
                xml = data.decode("utf-8")
                xml = re.sub(r"<w:commentRangeStart[^>]*/>", "", xml)
                xml = re.sub(r"<w:commentRangeEnd[^>]*/>", "", xml)
                xml = re.sub(r"<w:r>(?:(?!</w:r>).)*?<w:commentReference[^>]*/>.*?</w:r>", "", xml, flags=re.S)
                sect = re.search(r"<w:sectPr[ >].*?</w:sectPr>", xml, re.S).group(0)
                # The template already declares a first-page *footer*; look for a header one.
                if not re.search(r'<w:headerReference w:type="first"', sect):
                    xml = xml.replace(
                        '<w:headerReference w:type="default"',
                        f'<w:headerReference w:type="first" r:id="{rel_id}"/><w:headerReference w:type="default"', 1)
                data = xml.encode("utf-8")
            elif name == "word/header1.xml":
                xml = data.decode("utf-8")
                xml = xml.replace("<wp:posOffset>-635</wp:posOffset>", f"<wp:posOffset>{LOGO_LIFT_EMU}</wp:posOffset>")
                data = xml.encode("utf-8")
            elif name == "word/_rels/document.xml.rels":
                xml = data.decode("utf-8")
                xml = re.sub(r'<Relationship[^>]*Target="(' + "|".join(p[:-4] for p in COMMENT_PARTS) + r')\.xml"[^>]*/>',
                             "", xml)
                xml = xml.replace("</Relationships>",
                                  f'<Relationship Id="{rel_id}" Type="http://schemas.openxmlformats.org/'
                                  f'officeDocument/2006/relationships/header" Target="headerFirst.xml"/>'
                                  "</Relationships>")
                data = xml.encode("utf-8")
            elif name == "[Content_Types].xml":
                xml = data.decode("utf-8")
                for part in COMMENT_PARTS:
                    xml = re.sub(r'<Override PartName="/word/' + re.escape(part) + r'"[^>]*/>', "", xml)
                xml = xml.replace("</Types>",
                                  f'<Override PartName="/{first_header}" ContentType="application/vnd.'
                                  f'openxmlformats-officedocument.wordprocessingml.header+xml"/></Types>')
                data = xml.encode("utf-8")
            zout.writestr(item, data)
        zout.writestr(first_header, FIRST_PAGE_HEADER)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", help="The corporate Word template (.docx)")
    parser.add_argument("--out", default=str(Path(__file__).resolve().parents[1] / "templates" / "vertowave_base.docx"))
    args = parser.parse_args()

    src = Path(args.source)
    if not src.is_file():
        print(f"ERROR: {src} not found", file=sys.stderr)
        return 2
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".trimmed.docx")
    dropped = trim_to_cover(src, tmp)
    rewrite(tmp, out)
    tmp.unlink()

    doc = Document(str(out))
    print(f"dropped {dropped} blocks after the cover")
    print(f"wrote {out} ({out.stat().st_size} bytes), {len(doc.paragraphs)} cover blocks")
    print("first-page header:", "empty" if not doc.sections[0].first_page_header.paragraphs[0].text.strip() else "?")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
